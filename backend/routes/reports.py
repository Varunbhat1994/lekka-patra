"""PDF, Excel, and WhatsApp report routes."""
import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from core.database import db
from security.authentication import get_current_user
from services.ledger import compute_worker_ledger, compute_contractor_ledger


router = APIRouter()


def now_utc():
    return datetime.now(timezone.utc)


@router.get("/reports/pdf")
async def report_pdf(
    user: dict = Depends(get_current_user),
    worker_id: Optional[str] = None,
    start: Optional[str] = None,
    end: Optional[str] = None,
):
    workers_query = {"user_id": user["user_id"]}
    if worker_id:
        workers_query["id"] = worker_id
    workers = await db.workers.find(workers_query, {"_id": 0}).to_list(1000)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Farm Labor Report")
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Farm Labor Report", styles["Title"]))
    story.append(Paragraph(f"Owner: {user.get('name','')} · {user.get('email','')}", styles["Normal"]))
    story.append(Paragraph(f"Generated: {now_utc().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
    if start and end:
        story.append(Paragraph(f"Period: {start} to {end}", styles["Normal"]))
    story.append(Spacer(1, 12))

    for w in workers:
        led = await compute_worker_ledger(user["user_id"], w, start=start, end=end)
        story.append(Paragraph(f"<b>{w['name']}</b> ({w.get('skill','')}) — Rate: Rs {w['daily_rate']}", styles["Heading3"]))
        summary = [
            ["Days Worked", "Total Earned", "Advance", "Returned", "Settled", "Earned (period)", "Balance"],
            [led["days_worked"], f"Rs {led['total_earned']}",
             f"Rs {led['total_advance']}", f"Rs {led['total_returned']}",
             f"Rs {led.get('total_settled', 0)}", f"Rs {led['pending']}",
             f"Rs {led.get('final_balance', 0)}"],
        ]
        t = Table(summary, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2f6b3b")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("PADDING", (0,0), (-1,-1), 6),
        ]))
        story.append(t)

        # Date-wise attendance log
        if led["attendance"]:
            story.append(Spacer(1, 6))
            story.append(Paragraph("<b>Attendance (date-wise)</b>", styles["Normal"]))
            att_rows = [["Date", "Status", "OT hrs", "Field/Crop", "Description"]]
            for a in sorted(led["attendance"], key=lambda x: x["date"]):
                att_rows.append([
                    a["date"],
                    a["status"].replace("_", " ").title(),
                    a.get("overtime_hours", 0) or "",
                    a.get("field_crop", ""),
                    (a.get("description", "") or "")[:60],
                ])
            att_tab = Table(att_rows, hAlign="LEFT")
            att_tab.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eeeeee")),
                ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
                ("FONTSIZE", (0,0), (-1,-1), 9),
                ("PADDING", (0,0), (-1,-1), 4),
            ]))
            story.append(att_tab)

        if led["advances"] or led["returns"]:
            story.append(Spacer(1, 6))
            story.append(Paragraph("<b>Advances & Returns</b>", styles["Normal"]))
            rows = [["Date", "Type", "Amount", "Method", "Notes"]]
            for a in sorted(led["advances"], key=lambda x: x["date"]):
                rows.append([a["date"], "Advance", f"Rs {a['amount']}", a.get("method", ""), a.get("notes", "")])
            for r in sorted(led["returns"], key=lambda x: x["date"]):
                rows.append([r["date"], "Return", f"Rs {r['amount']}", r.get("method", ""), r.get("notes", "")])
            tab = Table(rows, hAlign="LEFT")
            tab.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eeeeee")),
                ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
                ("FONTSIZE", (0,0), (-1,-1), 9),
                ("PADDING", (0,0), (-1,-1), 4),
            ]))
            story.append(tab)
        story.append(Spacer(1, 16))

    doc.build(story)
    buf.seek(0)
    fname = "farm_report.pdf" if not worker_id else f"worker_{worker_id[:8]}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={fname}"})

@router.get("/reports/excel")
async def report_excel(
    user: dict = Depends(get_current_user),
    worker_id: Optional[str] = None,
    start: Optional[str] = None,
    end: Optional[str] = None,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Worker", "Skill", "Daily Rate", "Days Worked", "Total Earned", "Advance", "Returned", "Settled", "Earned (period)", "Balance"])
    workers_query = {"user_id": user["user_id"]}
    if worker_id:
        workers_query["id"] = worker_id
    workers = await db.workers.find(workers_query, {"_id": 0}).to_list(1000)
    for w in workers:
        led = await compute_worker_ledger(user["user_id"], w, start=start, end=end)
        ws.append([w["name"], w.get("skill",""), w["daily_rate"],
                   led["days_worked"], led["total_earned"],
                   led["total_advance"], led["total_returned"],
                   led.get("total_settled", 0), led["pending"],
                   led.get("final_balance", 0)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = "farm_report.xlsx" if not worker_id else f"worker_{worker_id[:8]}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"})

@router.get("/reports/contractor/{cid}/pdf")
async def contractor_pdf(
    cid: str,
    user: dict = Depends(get_current_user),
    start: Optional[str] = None,
    end: Optional[str] = None,
):
    contractor = await db.contractors.find_one({"id": cid, "user_id": user["user_id"]}, {"_id": 0})
    if not contractor:
        raise HTTPException(404, "Contractor not found")
    led = await compute_contractor_ledger(user["user_id"], contractor, start=start, end=end)
    visits = led["visits"]
    payments = led["payments"]
    returns = led["returns"]
    total_workers = led["total_workers_brought"]
    total_paid = led["total_paid"]
    total_returned = led["total_returned"]
    net_paid = led["net_paid"]
    total_settled = led.get("total_settled", 0)
    final_balance = float(led.get("final_balance", 0) or 0)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"Contractor · {contractor['name']}")
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Contractor Report — {contractor['name']}", styles["Title"]),
        Paragraph(f"Mobile: {contractor.get('mobile','—')}", styles["Normal"]),
        Paragraph(f"Generated: {now_utc().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
    ]
    if start and end:
        story.append(Paragraph(f"Period: {start} to {end}", styles["Normal"]))
    # Direction sentence
    if final_balance > 0:
        story.append(Paragraph(f"<b>You owe contractor Rs {round(final_balance, 2)}</b>", styles["Normal"]))
    elif final_balance < 0:
        story.append(Paragraph(f"<b>Contractor owes you Rs {round(-final_balance, 2)}</b>", styles["Normal"]))
    else:
        story.append(Paragraph("<b>Balanced</b>", styles["Normal"]))
    story.append(Spacer(1, 12))
    summary = [
        ["Visits", "Total Workers", "Total Paid", "Returned", "Settled", "Net Paid", "Balance"],
        [
            len(visits), total_workers,
            f"Rs {total_paid}", f"Rs {total_returned}",
            f"Rs {round(total_settled, 2)}", f"Rs {round(net_paid, 2)}",
            f"Rs {round(final_balance, 2)}",
        ],
    ]
    t = Table(summary, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2f6b3b")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
        ("PADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    if visits:
        story.append(Paragraph("<b>Visits</b>", styles["Heading4"]))
        rows = [["Date", "Workers", "Field/Crop", "Notes"]]
        for v in visits:
            rows.append([v["date"], v.get("workers_count", 0), v.get("field_crop",""), v.get("notes","")])
        vt = Table(rows, hAlign="LEFT")
        vt.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eeeeee")),
            ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
            ("FONTSIZE", (0,0), (-1,-1), 9),
            ("PADDING", (0,0), (-1,-1), 4),
        ]))
        story.append(vt)
        story.append(Spacer(1, 12))

    if payments or returns:
        story.append(Paragraph("<b>Payments & Returns</b>", styles["Heading4"]))
        rows = [["Date", "Type", "Amount", "Method", "Notes"]]
        for p in payments:
            rows.append([p["date"], "Payment", f"Rs {p['amount']}", p.get("method",""), p.get("notes","")])
        for r in returns:
            rows.append([r["date"], "Return", f"Rs {r['amount']}", r.get("method",""), r.get("notes","")])
        pt = Table(rows, hAlign="LEFT")
        pt.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eeeeee")),
            ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
            ("FONTSIZE", (0,0), (-1,-1), 9),
            ("PADDING", (0,0), (-1,-1), 4),
        ]))
        story.append(pt)

    doc.build(story)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=contractor_{cid[:8]}.pdf"})

@router.get("/reports/contractor/{cid}/excel")
async def contractor_excel(
    cid: str,
    user: dict = Depends(get_current_user),
    start: Optional[str] = None,
    end: Optional[str] = None,
):
    contractor = await db.contractors.find_one({"id": cid, "user_id": user["user_id"]}, {"_id": 0})
    if not contractor:
        raise HTTPException(404, "Contractor not found")
    led = await compute_contractor_ledger(user["user_id"], contractor, start=start, end=end)
    visits = led["visits"]
    payments = led["payments"]
    returns = led["returns"]

    wb = Workbook()
    s1 = wb.active
    s1.title = "Summary"
    s1.append(["Contractor", contractor["name"]])
    s1.append(["Mobile", contractor.get("mobile", "")])
    s1.append(["Visits", len(visits)])
    s1.append(["Total Workers Brought", sum(v.get("workers_count", 0) for v in visits)])
    s1.append(["Total Paid", sum(p["amount"] for p in payments)])
    s1.append(["Total Returned", sum(r["amount"] for r in returns)])
    s1.append(["Total Settled", led.get("total_settled", 0)])
    s1.append(["Net Paid", led.get("net_paid", 0)])
    s1.append(["Balance", led.get("final_balance", 0)])
    _fb = float(led.get("final_balance", 0) or 0)
    if _fb > 0:
        s1.append(["Direction", f"You owe contractor Rs {round(_fb, 2)}"])
    elif _fb < 0:
        s1.append(["Direction", f"Contractor owes you Rs {round(-_fb, 2)}"])
    else:
        s1.append(["Direction", "Balanced"])

    s2 = wb.create_sheet("Visits")
    s2.append(["Date", "Workers", "Field/Crop", "Notes"])
    for v in visits:
        s2.append([v["date"], v.get("workers_count", 0), v.get("field_crop",""), v.get("notes","")])

    s3 = wb.create_sheet("Payments")
    s3.append(["Date", "Amount", "Method", "Notes"])
    for p in payments:
        s3.append([p["date"], p["amount"], p.get("method",""), p.get("notes","")])

    s4 = wb.create_sheet("Returns")
    s4.append(["Date", "Amount", "Method", "Notes"])
    for r in returns:
        s4.append([r["date"], r["amount"], r.get("method",""), r.get("notes","")])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=contractor_{cid[:8]}.xlsx"})

@router.get("/reports/whatsapp/{worker_id}")
async def whatsapp_text(worker_id: str, user: dict = Depends(get_current_user), lang: str = "en"):
    worker = await db.workers.find_one({"id": worker_id, "user_id": user["user_id"]}, {"_id": 0})
    if not worker:
        raise HTTPException(404, "Worker not found")
    led = await compute_worker_ledger(user["user_id"], worker)
    fb = float(led.get("final_balance", 0) or 0)
    if lang == "kn":
        if fb > 0:
            bal_line = f"ಬಾಕಿ: ನೀವು ಕಾರ್ಮಿಕರಿಗೆ ರೂ {round(fb, 2)} ಸಾಲ"
        elif fb < 0:
            bal_line = f"ಬಾಕಿ: ಕಾರ್ಮಿಕ ನಿಮಗೆ ರೂ {round(-fb, 2)} ಸಾಲ"
        else:
            bal_line = "ಬಾಕಿ: ಸಮತೋಲನ"
        msg = (
            f"ನಮಸ್ಕಾರ {worker['name']},\n"
            f"ಒಟ್ಟು ಕೆಲಸದ ದಿನಗಳು: {led['days_worked']}\n"
            f"ಒಟ್ಟು ಸಂಬಳ: ರೂ {led['total_earned']}\n"
            f"ಮುಂಗಡ ಪಾವತಿ: ರೂ {led['total_advance']}\n"
            f"{bal_line}"
        )
    else:
        if fb > 0:
            bal_line = f"Balance: You owe worker Rs {round(fb, 2)}"
        elif fb < 0:
            bal_line = f"Balance: Worker owes you Rs {round(-fb, 2)}"
        else:
            bal_line = "Balance: Balanced"
        msg = (
            f"Hi {worker['name']},\n"
            f"Days Worked: {led['days_worked']}\n"
            f"Total Earned: Rs {led['total_earned']}\n"
            f"Advance Paid: Rs {led['total_advance']}\n"
            f"{bal_line}"
        )
    return {"message": msg, "phone": worker.get("mobile", "")}

@router.get("/reports/contractor/{cid}/whatsapp")
async def contractor_whatsapp(cid: str, user: dict = Depends(get_current_user), lang: str = "en"):
    contractor = await db.contractors.find_one({"id": cid, "user_id": user["user_id"]}, {"_id": 0})
    if not contractor:
        raise HTTPException(404, "Contractor not found")
    visits = await db.contractor_visits.find(
        {"contractor_id": cid, "user_id": user["user_id"]}, {"_id": 0}
    ).sort("date", -1).to_list(500)
    payments = await db.contractor_payments.find(
        {"contractor_id": cid, "user_id": user["user_id"]}, {"_id": 0}
    ).sort("date", -1).to_list(500)
    returns = await db.contractor_returns.find(
        {"contractor_id": cid, "user_id": user["user_id"]}, {"_id": 0}
    ).sort("date", -1).to_list(500)
    total_workers = sum(v.get("workers_count", 0) for v in visits)
    total_paid = sum(p["amount"] for p in payments)
    total_returned = sum(r["amount"] for r in returns)
    net_paid = round(total_paid - total_returned, 2)
    # Compute contractor-scoped settled + direction via the same ledger.
    led = await compute_contractor_ledger(user["user_id"], contractor)
    total_settled = led.get("total_settled", 0)
    fb = float(led.get("final_balance", 0) or 0)

    recent_visits = visits[:5]
    recent_pays = payments[:5]

    if lang == "kn":
        if fb > 0:
            bal_line = f"ಬಾಕಿ: ನೀವು ಗುತ್ತಿಗೆದಾರರಿಗೆ ರೂ {round(fb, 2)} ಸಾಲ"
        elif fb < 0:
            bal_line = f"ಬಾಕಿ: ಗುತ್ತಿಗೆದಾರ ನಿಮಗೆ ರೂ {round(-fb, 2)} ಸಾಲ"
        else:
            bal_line = "ಬಾಕಿ: ಸಮತೋಲನ"
        lines = [
            f"ನಮಸ್ಕಾರ {contractor['name']},",
            "",
            f"ಒಟ್ಟು ಭೇಟಿಗಳು: {len(visits)}",
            f"ಒಟ್ಟು ಕಾರ್ಮಿಕರು: {total_workers}",
            f"ಒಟ್ಟು ಪಾವತಿ: ರೂ {total_paid}",
            f"ವಾಪಸಾತಿ: ರೂ {total_returned}",
            f"ಒಟ್ಟು ಇತ್ಯರ್ಥ: ರೂ {total_settled}",
            f"ನಿವ್ವಳ ಪಾವತಿ: ರೂ {net_paid}",
            bal_line,
        ]
        if recent_visits:
            lines += ["", "ಇತ್ತೀಚಿನ ಭೇಟಿಗಳು:"]
            for v in recent_visits:
                lines.append(f"• {v['date']} — {v.get('workers_count', 0)} ಕಾರ್ಮಿಕರು ({v.get('field_crop','—')})")
        if recent_pays:
            lines += ["", "ಇತ್ತೀಚಿನ ಪಾವತಿಗಳು:"]
            for p in recent_pays:
                lines.append(f"• {p['date']} — ರೂ {p['amount']} ({p.get('method','')})")
    else:
        if fb > 0:
            bal_line = f"Balance: You owe contractor Rs {round(fb, 2)}"
        elif fb < 0:
            bal_line = f"Balance: Contractor owes you Rs {round(-fb, 2)}"
        else:
            bal_line = "Balance: Balanced"
        lines = [
            f"Hi {contractor['name']},",
            "",
            f"Total visits: {len(visits)}",
            f"Total workers brought: {total_workers}",
            f"Total paid: Rs {total_paid}",
            f"Returned: Rs {total_returned}",
            f"Total settled: Rs {total_settled}",
            f"Net paid: Rs {net_paid}",
            bal_line,
        ]
        if recent_visits:
            lines += ["", "Recent visits:"]
            for v in recent_visits:
                lines.append(f"• {v['date']} — {v.get('workers_count', 0)} workers ({v.get('field_crop','—')})")
        if recent_pays:
            lines += ["", "Recent payments:"]
            for p in recent_pays:
                lines.append(f"• {p['date']} — Rs {p['amount']} ({p.get('method','')})")
    return {"message": "\n".join(lines), "phone": contractor.get("mobile", "")}
