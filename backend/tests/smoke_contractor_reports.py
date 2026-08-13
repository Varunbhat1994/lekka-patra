"""Live HTTP smoke test for contractor PDF/Excel/WhatsApp exports.

Verifies:
- /api/contractors/{id}/ledger returns total_settled + final_balance
- /api/reports/contractor/{id}/pdf renders (real bytes) and its extracted
  text contains "Balance" and the direction sentence
- /api/reports/contractor/{id}/excel Summary sheet contains the new rows
- /api/reports/contractor/{id}/whatsapp text includes the balance line
- After Mark Settled, PDF text reads "Balanced" and ledger totals settle
"""
import asyncio
import io
import os
import sys
import uuid
from datetime import datetime, timezone, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import httpx
from core.database import db

BASE = os.environ.get(
    "REACT_APP_BACKEND_URL",
    "https://field-crew-log-1.preview.emergentagent.com",
) + "/api"


async def _seed_session(uid: str, mobile: str) -> str:
    now = datetime.now(timezone.utc)
    await db.users.insert_one({
        "user_id": uid, "email": f"{uid}@t.com", "mobile": mobile,
        "name": "CSmoke", "role": "owner",
        "subscription_active": True,
        "subscription_expires_at": (now + timedelta(days=365)).isoformat(),
        "trial_starts_at": now.isoformat(),
        "trial_expires_at": (now + timedelta(days=365)).isoformat(),
        "created_at": now.isoformat(),
    })
    token = uuid.uuid4().hex + uuid.uuid4().hex
    await db.user_sessions.insert_one({
        "session_token": token, "user_id": uid,
        "expires_at": (now + timedelta(hours=2)).isoformat(),
        "created_at": now.isoformat(),
    })
    return token


def _extract_pdf_text(pdf_bytes: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(p.extract_text() or "" for p in reader.pages)


async def main() -> None:
    uid = f"csm_{uuid.uuid4().hex[:8]}"
    token = await _seed_session(uid, "9000000200")
    h = {"Authorization": f"Bearer {token}"}
    try:
        async with httpx.AsyncClient(base_url=BASE, headers=h, timeout=60, verify=False) as c:
            cont = (await c.post("/contractors", json={"name": "CSmoke", "mobile": "", "notes": ""})).json()
            await c.post("/contractor-payments", json={"contractor_id": cont["id"], "date": "2026-01-10", "amount": 5000, "method": "cash"})
            await c.post("/contractor-returns",  json={"contractor_id": cont["id"], "date": "2026-01-20", "amount": 2000, "method": "cash"})

            # Ledger
            led = (await c.get(f"/contractors/{cont['id']}/ledger")).json()
            assert led["net_paid"] == 3000.0 and led["final_balance"] == -3000.0 and led["total_settled"] == 0.0, led
            print(f"✓ /contractors/{{id}}/ledger  final_balance={led['final_balance']}, total_settled={led['total_settled']}")

            # PDF
            pdf = await c.get(f"/reports/contractor/{cont['id']}/pdf")
            assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
            text = _extract_pdf_text(pdf.content)
            assert "Balance" in text and "Contractor owes you" in text and "3000" in text, text[:500]
            print(f"✓ PDF ({len(pdf.content)} bytes) — text: 'Contractor owes you Rs 3000.0'")

            # Excel
            xls = await c.get(f"/reports/contractor/{cont['id']}/excel")
            assert xls.status_code == 200 and len(xls.content) > 1000
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(xls.content))
            labels = [row[0] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]]
            for want in ("Total Paid", "Total Returned", "Total Settled", "Net Paid", "Balance", "Direction"):
                assert want in labels, f"Missing {want} in Excel: {labels}"
            print("✓ Excel Summary sheet includes: Total Paid, Total Returned, Total Settled, Net Paid, Balance, Direction")

            # WhatsApp
            wa = (await c.get(f"/reports/contractor/{cont['id']}/whatsapp?lang=en")).json()
            assert "Contractor owes you" in wa["message"], wa
            bal_line = [ln for ln in wa["message"].splitlines() if "Balance" in ln][0]
            print(f"✓ WhatsApp: '{bal_line}'")

            # Mark Settled → check post-settle state
            await c.post("/settlements", json={"contractor_id": cont["id"], "up_to_date": "2026-02-01"})
            led2 = (await c.get(f"/contractors/{cont['id']}/ledger")).json()
            assert led2["final_balance"] == 0.0 and led2["total_settled"] == 3000.0, led2
            pdf2 = await c.get(f"/reports/contractor/{cont['id']}/pdf")
            text2 = _extract_pdf_text(pdf2.content)
            assert "Balanced" in text2, text2[:500]
            print(f"✓ After Mark Settled: final_balance=0, total_settled=3000; PDF shows 'Balanced'")

            print("\nALL CONTRACTOR REPORT SMOKE CHECKS PASSED")
    finally:
        for coll in ("users", "user_sessions", "contractors",
                     "contractor_payments", "contractor_returns", "settlements"):
            await db[coll].delete_many({"user_id": uid})
        await db.user_sessions.delete_many({"session_token": token})


if __name__ == "__main__":
    asyncio.run(main())
